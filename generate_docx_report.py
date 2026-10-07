"""Generate the Vietnamese system architecture document (.docx).

The document content is declared once in :data:`CONTENT` and rendered by two
writers: :func:`build_docx` (Word file) and :func:`render_markdown` (raw
text). Using one source guarantees the Word file and the plain-text version
never diverge.

Usage::

    python generate_docx_report.py                       # write .docx
    python generate_docx_report.py --markdown out.md     # also dump text
"""

from __future__ import annotations

import argparse
import datetime as _dt
import sys
from pathlib import Path
from typing import List, Optional, Sequence, Tuple, Union

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT_NAME: str = "Kien_truc_He_thong_Aeon_Fall_Detection.docx"
TITLE: str = "TÀI LIỆU KIẾN TRÚC HỆ THỐNG"
SUBTITLE: str = "Hệ thống phát hiện té ngã Zero-shot Aeon_fall_detection"
VERSION: str = "1.0"

#: Block types: ("h1"|"h2", text), ("p", text), ("bullets", [items]),
#: ("table", [header], [[row], ...]).
Block = Union[Tuple[str, str], Tuple[str, List[str]],
              Tuple[str, List[str], List[List[str]]]]

CONTENT: List[Block] = [
    ("h1", "1. Tổng quan hệ thống"),
    ("h2", "1.1. Mục tiêu"),
    ("p", "Hệ thống phát hiện té ngã theo thời gian thực từ luồng camera, "
          "không sử dụng dữ liệu huấn luyện chuyên biệt về té ngã "
          "(Zero-shot). Té ngã được suy luận từ quan hệ ngữ nghĩa giữa "
          "người và mặt sàn trong Đồ thị cảnh (Scene Graph)."),
    ("h2", "1.2. Nguyên lý"),
    ("bullets", [
        "Phát hiện đối tượng từ vựng mở: YOLO-World v2 (kiến trúc YOLOv8) "
        "định vị “person”, “floor” và các bề mặt bổ sung theo mô tả văn "
        "bản.",
        "Dự đoán quan hệ: RelateAnything chấm điểm bộ ba (chủ thể, vị từ, "
        "đối tượng), ví dụ “person – lying on – floor”, với từ vựng vị từ "
        "thay đổi khi chạy.",
        "Suy luận không gian – thời gian: kết hợp điểm quan hệ với đặc "
        "trưng hình học, lọc theo cửa sổ thời gian trên từng đối tượng "
        "được theo vết.",
    ]),
    ("h2", "1.3. Đặc tính chính"),
    ("table", ["Đặc tính", "Mô tả"], [
        ["Zero-shot", "Không cần dữ liệu té ngã; mở rộng bằng từ vựng văn "
                      "bản."],
        ["Khả giải thích", "Mỗi cảnh báo kèm bộ ba quan hệ và điểm số."],
        ["Đa nền tảng", "Cùng một API cho CPU Biên (ONNX/OpenVINO) và GPU "
                        "máy chủ (PyTorch/CUDA)."],
        ["Cấu hình động", "Thay đổi từ vựng, ngưỡng khi đang chạy, không "
                          "huấn luyện lại."],
    ]),

    ("h1", "2. Thiết kế kiến trúc"),
    ("h2", "2.1. Thành phần"),
    ("table", ["Thành phần", "Tệp", "Chức năng"], [
        ["Khởi tạo môi trường", "setup.sh",
         "Tạo môi trường conda, tải RelateAnything, cài đặt phụ thuộc."],
        ["Tối ưu hóa mô hình", "export_onnx.py",
         "Xuất ONNX với trục động, ngân hàng vị từ, IR OpenVINO FP16, "
         "kiểm định tương đương."],
        ["Bộ phát hiện", "pipeline.py",
         "YOLO-World + ByteTrack; xuất OpenVINO theo từ vựng, lưu đệm."],
        ["Backend quan hệ", "pipeline.py",
         "TORCH_GPU (PyTorch/CUDA) hoặc ONNX_CPU (OpenVINO/ONNX Runtime)."],
        ["Bộ suy luận", "pipeline.py",
         "Hợp nhất bằng chứng, lọc thời gian, sinh cảnh báo."],
        ["Giao diện", "app.py",
         "Gradio: video/webcam, từ vựng, ngưỡng, nhật ký cảnh báo."],
        ["Đóng gói", "Dockerfile, docker-compose.yml",
         "Ảnh CUDA 12.9, cấp phát GPU NVIDIA."],
    ]),
    ("h2", "2.2. Luồng xử lý"),
    ("bullets", [
        "Bước 1 – Đọc khung hình: nhận khung BGR từ webcam, RTSP hoặc tệp "
        "video.",
        "Bước 2 – Phát hiện đối tượng: YOLO-World trả về hộp bao, nhãn, "
        "mã theo vết.",
        "Bước 3 – Chọn vùng: ưu tiên người, sàn, bề mặt an toàn; tối đa 16 "
        "hộp. Thiếu sàn thì dùng sàn ảo (25% đáy khung hình).",
        "Bước 4 – Dự đoán quan hệ: RelateAnything trích xuất đặc trưng "
        "DINOv3, trả về logit vị từ [K, V] và logit cặp [K].",
        "Bước 5 – Tính quan hệ không gian: điểm = σ(a·(logit vị từ + logit "
        "cặp) + b); bằng chứng quan hệ = té / (té + đứng) trên cặp "
        "người→sàn.",
        "Bước 6 – Hợp nhất: điểm = 0,6·quan hệ + 0,4·hình học (tư thế, "
        "tiếp xúc sàn, độ sụt chiều cao).",
        "Bước 7 – Lọc thời gian: cửa sổ 8 khung, tối thiểu 4 khung vượt "
        "ngưỡng thì xác nhận TÉ NGÃ; có trễ (hysteresis) để chống dao động.",
        "Bước 8 – Đầu ra: khung hình chú thích, trạng thái, cảnh báo.",
    ]),
    ("h2", "2.3. Quy tắc quyết định"),
    ("table", ["Tín hiệu", "Định nghĩa", "Vai trò"], [
        ["Quan hệ", "max(té) / (max(té) + max(đứng))",
         "Khử độ lệch điểm theo cảnh."],
        ["Tư thế", "clip((w/h − 0,6) / 0,8)", "Hộp nằm ngang khi nằm."],
        ["Tiếp xúc sàn", "IoA(người, sàn)", "Phân biệt sàn và đồ vật."],
        ["Sụt chiều cao", "1 − h hiện tại / h lớn nhất",
         "Bắt quá trình ngã."],
        ["Bề mặt an toàn", "Nằm trên giường/sofa ⇒ điểm × 0,3",
         "Loại trừ nghỉ ngơi bình thường."],
    ]),

    ("h1", "3. Tối ưu hóa phần cứng (Edge CPU vs Server GPU)"),
    ("h2", "3.1. Biên (Edge): Intel Core i7-1260P"),
    ("bullets", [
        "Chế độ InferenceMode.ONNX_CPU; không cần PyTorch khi chạy.",
        "Mô hình quan hệ: IR OpenVINO, trọng số FP16, tính toán FP32, "
        "PERFORMANCE_HINT=LATENCY ưu tiên lõi P.",
        "Bộ phát hiện: xuất OpenVINO một lần theo từ vựng; nhúng sẵn đặc "
        "trưng văn bản CLIP.",
        "Hình dạng tĩnh: 16 khe hộp, đệm 0 kèm box_counts; từ vựng V là "
        "trục động.",
        "Dự phòng: ONNX Runtime CPU khi thiếu IR.",
        "Kết quả đo trên i7-1260P: mô hình quan hệ 259 ms (OpenVINO), "
        "276 ms (ONNX Runtime); bộ phát hiện 30–40 ms; toàn luồng 2–3 "
        "khung hình/giây.",
    ]),
    ("h2", "3.2. Máy chủ: NVIDIA RTX 4090 24 GB / RTX 5090 32 GB"),
    ("bullets", [
        "Chế độ InferenceMode.TORCH_GPU; mã hóa vị từ tùy ý bằng bộ mã hóa "
        "văn bản.",
        "TF32 cho phép nhân ma trận; bộ phát hiện chạy FP16.",
        "CUDA 12.9: hỗ trợ đồng thời Ada (sm_89) và Blackwell (sm_120).",
        "Mỗi luồng camera một engine; ước tính 2–3 GB VRAM/engine.",
        "Mở rộng: TensorRT từ cùng đồ thị ONNX.",
    ]),
    ("h2", "3.3. So sánh"),
    ("table",
     ["Tiêu chí", "Edge CPU (i7-1260P)", "Server GPU (RTX 4090/5090)"],
     [
         ["Backend", "OpenVINO / ONNX Runtime", "PyTorch CUDA"],
         ["Độ chính xác số", "FP32 (trọng số FP16)", "TF32 / FP16"],
         ["Từ vựng vị từ", "Trong ngân hàng vị từ", "Tùy ý"],
         ["Số luồng camera", "1", "Nhiều (theo VRAM)"],
         ["Mục đích", "Thử nghiệm, triển khai tại chỗ",
          "Vận hành tập trung"],
     ]),

    ("h1", "4. Quy trình triển khai (Deployment)"),
    ("h2", "4.1. Chuẩn bị"),
    ("bullets", [
        "Chạy setup.sh: tạo môi trường Aeon_fall_detection (Python 3.12, "
        "yêu cầu tối thiểu của RelateAnything), tự chọn gói CUDA 12.9 hoặc "
        "CPU.",
        "Chạy export_onnx.py --openvino: sinh ONNX, IR OpenVINO, ngân hàng "
        "vị từ; kiểm định tương đương với PyTorch.",
        "Sao chép thư mục models/relsgg-vits16plus tới thiết bị Biên.",
    ]),
    ("h2", "4.2. Triển khai"),
    ("table", ["Môi trường", "Lệnh", "Chế độ"], [
        ["Biên / máy trạm", "python app.py --mode onnx_cpu", "ONNX_CPU"],
        ["Máy chủ GPU", "python app.py --mode torch_gpu --host 0.0.0.0",
         "TORCH_GPU"],
        ["Container GPU", "docker compose up -d --build", "TORCH_GPU"],
        ["Xuất mô hình", "docker compose --profile export run --rm export",
         "—"],
    ]),
    ("h2", "4.3. Container"),
    ("bullets", [
        "Ảnh nền nvidia/cuda:12.9.1-cudnn-runtime-ubuntu24.04; driver "
        "NVIDIA ≥ 575, NVIDIA Container Toolkit.",
        "Cấp phát GPU qua deploy.resources.reservations.devices và "
        "runtime: nvidia.",
        "Volume /data lưu bộ đệm Hugging Face, trọng số, mô hình xuất.",
        "HEALTHCHECK cổng 7860; --preload dừng container nếu GPU/mô hình "
        "lỗi; chạy người dùng không đặc quyền.",
    ]),
    ("h2", "4.4. Vận hành"),
    ("bullets", [
        "Phát hành blue/green theo thẻ ảnh; kiểm định trên video mẫu trước "
        "khi chuyển lưu lượng.",
        "Giám sát độ trễ từng tầng, tần suất cảnh báo, tỉ lệ dùng sàn ảo.",
        "Thứ tự hiệu chỉnh: fall_threshold → min_fall_frames → "
        "relation_weight.",
        "Giấy phép: RelateAnything và Ultralytics theo AGPL-3.0; dịch vụ "
        "mạng phải tuân thủ hoặc mua giấy phép thương mại.",
    ]),
]


def _set_font(run_or_style: object, size: Optional[float] = None,
              bold: Optional[bool] = None,
              color: Optional[RGBColor] = None) -> None:
    """Apply Times New Roman (incl. East-Asian slot) to a run or style.

    Word picks the font for Vietnamese diacritics from the ``eastAsia``/
    ``cs`` slots in some locales; setting only ``font.name`` can yield mixed
    fonts inside one word, so all slots are set explicitly.

    Args:
        run_or_style: ``docx`` run or style object.
        size: Font size in points.
        bold: Bold flag.
        color: Font colour.
    """
    font = run_or_style.font
    font.name = "Times New Roman"
    element = (run_or_style.element if hasattr(run_or_style, "element")
               else run_or_style._element)
    rpr = element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(slot), "Times New Roman")
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color


def _shade(cell: object, hex_fill: str) -> None:
    """Fill a table cell background.

    Args:
        cell: ``docx`` table cell.
        hex_fill: RGB hex string such as ``"1F3864"``.
    """
    tcpr = cell._element.get_or_add_tcPr()
    shd = tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear",
                                         qn("w:color"): "auto",
                                         qn("w:fill"): hex_fill})
    tcpr.append(shd)


def build_docx(path: Path, blocks: Sequence[Block]) -> Path:
    """Render the content into a formatted Word document.

    Args:
        path: Output ``.docx`` path.
        blocks: Content blocks.

    Returns:
        Path: The written file.

    Raises:
        ValueError: If a block type is unknown.
        OSError: If the file cannot be written (e.g. open in Word).
    """
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)

    _set_font(doc.styles["Normal"], size=13)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.3
    for name, size in (("Heading 1", 15), ("Heading 2", 13.5)):
        _set_font(doc.styles[name], size=size, bold=True,
                  color=RGBColor(0x1F, 0x38, 0x64))

    # Title block.
    for text, size, bold in ((TITLE, 18, True), (SUBTITLE, 14, True)):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(para.add_run(text), size=size, bold=bold,
                  color=RGBColor(0x1F, 0x38, 0x64))
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(meta.add_run(f"Phiên bản {VERSION} – "
                           f"{_dt.date.today():%d/%m/%Y}"), size=11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(footer.add_run("Aeon_fall_detection – Tài liệu kiến trúc hệ "
                             "thống"), size=9)

    for block in blocks:
        kind = block[0]
        if kind in ("h1", "h2"):
            doc.add_heading(block[1], level=1 if kind == "h1" else 2)
        elif kind == "p":
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _set_font(para.add_run(block[1]))
        elif kind == "bullets":
            for item in block[1]:
                para = doc.add_paragraph(style="List Bullet")
                _set_font(para.add_run(item))
        elif kind == "table":
            header, rows = block[1], block[2]
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, text in zip(table.rows[0].cells, header):
                cell.text = ""
                _set_font(cell.paragraphs[0].add_run(text), size=12,
                          bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                _shade(cell, "1F3864")
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = ""
                    _set_font(cell.paragraphs[0].add_run(text), size=12)
            doc.add_paragraph()
        else:
            raise ValueError(f"Unknown block type: {kind}")
    doc.save(str(path))
    return path


def render_markdown(blocks: Sequence[Block]) -> str:
    """Render the same content as Markdown text.

    Args:
        blocks: Content blocks.

    Returns:
        str: Markdown document.
    """
    lines: List[str] = [f"# {TITLE}", "", f"**{SUBTITLE}**", "",
                        f"Phiên bản {VERSION}", ""]
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            lines += [f"## {block[1]}", ""]
        elif kind == "h2":
            lines += [f"### {block[1]}", ""]
        elif kind == "p":
            lines += [block[1], ""]
        elif kind == "bullets":
            lines += [f"- {item}" for item in block[1]] + [""]
        elif kind == "table":
            header, rows = block[1], block[2]
            lines.append("| " + " | ".join(header) + " |")
            lines.append("|" + "---|" * len(header))
            lines += ["| " + " | ".join(r) + " |" for r in rows] + [""]
    return "\n".join(lines).rstrip() + "\n"


def main(argv: Optional[Sequence[str]] = None) -> int:
    """Write the .docx (and optionally the Markdown text).

    Args:
        argv: Argument list (defaults to ``sys.argv``).

    Returns:
        int: Exit code.
    """
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--out-dir", default=".")
    parser.add_argument("--markdown", default="",
                        help="Also write the raw text as Markdown here.")
    args = parser.parse_args(argv)
    try:
        out_dir = Path(args.out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        path = build_docx(out_dir / OUTPUT_NAME, CONTENT)
        print(f"Wrote {path.resolve()}")
        if args.markdown:
            Path(args.markdown).write_text(render_markdown(CONTENT),
                                           encoding="utf-8")
            print(f"Wrote {Path(args.markdown).resolve()}")
        return 0
    except (OSError, ValueError) as exc:
        # The most common OSError is the file being open in Word (locked).
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
